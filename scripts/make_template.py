"""Maakt templates/RFIG-importtemplate.xlsx (één tabblad per importsoort + uitleg)."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "templates" / "RFIG-importtemplate.xlsx"
HEAD = PatternFill("solid", fgColor="1F5F5B")
WHITE = Font(color="FFFFFF", bold=True)

SHEETS = {
    "Producten": (["ean", "sku", "titel", "variant", "productgroep", "type", "eigenaar", "tags"],
                  "type: SEIZOEN of DOORLOPEND"),
    "Kosten": (["ean", "kostensoort", "bedrag_per_stuk_ex_btw", "geldig_vanaf", "geldig_tot", "basis"],
               "kostensoort: KOSTPRIJS, TRANSPORT, INVOERRECHTEN, FULFILMENT, VERZENDING, VERPAKKING, OPSLAG, RETOUR, OVERIG · basis: GEBOEKT of GESCHAT"),
    "Voorraad extern": (["ean", "locatie", "aantal", "peildatum"], "locatie bv. PACKFEED"),
    "Inkomend": (["ean", "aantal", "verwachte_datum", "bevestigd", "opmerking"], "bevestigd: JA of NEE (alleen bevestigde leveringen tellen mee in prognoses)"),
    "Doelen": (["ean_of_groep", "doeldatum", "eindvoorraad", "minimumprijs", "maximumprijs", "min_bijdrage_per_stuk", "prijsstap", "min_meetdagen", "marge_dagen"],
               "ean_of_groep: EAN of GROEP:<productgroep>, bv. GROEP:Dekens"),
}

wb = Workbook()
info = wb.active
info.title = "Uitleg"
lines = [
    "RFIG Dashboard — importtemplate",
    "",
    "Vul per tabblad één regel per record in. Laat de kopregel staan.",
    "EAN-kolommen staan als tekst opgemaakt: voorloopnullen blijven behouden.",
    "Bedragen excl. btw, per stuk. Datums als JJJJ-MM-DD (of DD-MM-JJJJ).",
    "Ontbrekende kosten niet als 0 invullen: leeg laten. Het dashboard toont dan 'onvolledig'.",
    "Importeren: dashboard → Gegevens & instellingen → Import. Eerst proefimport, dan vastleggen.",
    "",
    "Toelichting per tabblad:",
] + [f"- {k}: {v[1]}" for k, v in SHEETS.items()]
for i, t in enumerate(lines, 1):
    info.cell(row=i, column=1, value=t)
info["A1"].font = Font(bold=True, size=14)
info.column_dimensions["A"].width = 120

for name, (cols, note) in SHEETS.items():
    ws = wb.create_sheet(name)
    for c, h in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(horizontal="left")
        ws.column_dimensions[cell.column_letter].width = max(14, len(h) + 4)
    ws.freeze_panes = "A2"
    # EAN-kolom als tekst
    for r in range(2, 502):
        ws.cell(row=r, column=1).number_format = "@"
    ws.cell(row=1, column=len(cols) + 2, value=note).font = Font(italic=True, color="666666")
    if name == "Kosten":
        dv = DataValidation(type="list", formula1='"KOSTPRIJS,TRANSPORT,INVOERRECHTEN,FULFILMENT,VERZENDING,VERPAKKING,OPSLAG,RETOUR,OVERIG"', allow_blank=True)
        ws.add_data_validation(dv); dv.add("B2:B500")
        dv2 = DataValidation(type="list", formula1='"GEBOEKT,GESCHAT"', allow_blank=True)
        ws.add_data_validation(dv2); dv2.add("F2:F500")
    if name == "Producten":
        dv = DataValidation(type="list", formula1='"SEIZOEN,DOORLOPEND"', allow_blank=True)
        ws.add_data_validation(dv); dv.add("F2:F500")
    if name == "Inkomend":
        dv = DataValidation(type="list", formula1='"JA,NEE"', allow_blank=True)
        ws.add_data_validation(dv); dv.add("D2:D500")

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(OUT)
